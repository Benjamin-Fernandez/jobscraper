"""Read the company workbook. Read-only: the program never writes back to it."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

WANTED = {
    "company": "name",
    "tier": "tier",
    "category": "category",
    "careers page": "careers_url",
    "role type": "role_type_hint",
}


def read_companies(path: Path, sheet: str) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(
            f"sheet {sheet!r} not found in {path.name}; have: {wb.sheetnames}")
    ws = wb[sheet]

    header_map: dict[int, str] = {}
    rows: list[dict[str, Any]] = []
    ordinal = 0

    for row in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in row]
        if not header_map:
            lowered = [c.lower() for c in cells]
            if "company" in lowered:
                for idx, name in enumerate(lowered):
                    if name in WANTED:
                        header_map[idx] = WANTED[name]
            continue

        rec = {v: "" for v in WANTED.values()}
        for idx, key in header_map.items():
            if idx < len(cells):
                rec[key] = cells[idx]
        if not rec["name"]:
            continue
        ordinal += 1
        rec["ordinal"] = ordinal
        rows.append(rec)

    wb.close()
    if not rows:
        raise SystemExit(f"no company rows found in {path.name}:{sheet}")
    return rows
