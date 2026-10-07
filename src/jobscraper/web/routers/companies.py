"""`/api/companies/list` - the user's own list of companies (M16).

`PUT` takes the list as a file - TXT or DOCX, one company per line, in the
user's order of preference - as the raw request body, like the resume upload
(M11-T4). A form upload is refused: a cross-site page cannot send a raw PUT
without a CORS preflight the app never grants, and the write guard refuses a
foreign Origin as well. The names are parsed here (models.parse_company_lines)
and stored; a name searched before keeps its result.

`GET` answers what the Companies tab shows: every company in order with its
search status, the counts, and the plan's cap on companies. The search itself
is a background job (`POST /api/jobs/companies`), never run here - the web
layer does not reach a model or the network (PRD 8.2).
"""
from __future__ import annotations

import csv
import datetime as _dt
import io
import zipfile
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import Response

from jobscraper import watchlist
from jobscraper.config import Config
from jobscraper.models import MAX_COMPANY_LINES, parse_company_lines
from jobscraper.web.control import companies_view, company_health_view
from jobscraper.web.deps import get_config, open_store

router = APIRouter(tags=["companies"])

MAX_BYTES = 1024 * 1024
TXT_TYPE = "text/plain"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@router.get("/companies/list")
def get_list(request: Request, cfg: Config = Depends(get_config)) -> dict[str, Any]:
    """The list in order, each company's search status, counts and the cap."""
    with open_store(request) as store:
        return companies_view(cfg, store)


@router.put("/companies/list")
async def put_list(request: Request, cfg: Config = Depends(get_config)) -> dict[str, Any]:
    """Replace the list with the uploaded file's companies, in its order."""
    media = request.headers.get("content-type", "").split(";", 1)[0].strip().lower()
    if media not in (TXT_TYPE, DOCX_TYPE):
        raise HTTPException(
            status_code=415,
            detail=f"send the file itself as the body, with Content-Type {TXT_TYPE} "
                   f"or {DOCX_TYPE} (not a form upload)")
    body = bytearray()
    async for chunk in request.stream():
        body.extend(chunk)
        if len(body) > MAX_BYTES:
            raise HTTPException(status_code=413, detail="the list is over 1 MB")
    text = _text(media, bytes(body))
    names = parse_company_lines(text)
    if not names:
        raise HTTPException(status_code=422,
                            detail="no company names found - put one company on each line")
    with open_store(request) as store:
        before = _found_keys(store)
        change = store.replace_company_list(names)
        _unwatch_dropped(cfg, store, before)
        view = companies_view(cfg, store)
    view["upload"] = dict(change, read=len(names), capped=len(names) >= MAX_COMPANY_LINES)
    return view


@router.get("/companies/health")
def get_health(request: Request) -> dict[str, Any]:
    """Every watched company and whether it can be scanned (M17)."""
    with open_store(request) as store:
        return company_health_view(store)


CSV_COLUMNS = [("company", "Company"), ("state", "Can be scanned"), ("reason", "Why not"),
               ("open_roles", "Open roles at last scan"), ("job_board", "Job board"),
               ("careers_url", "Careers URL"), ("last_scan", "Last scanned (UTC)"),
               ("last_success", "Last successful scan (UTC)"),
               ("failures_in_a_row", "Failures in a row"), ("error", "Error detail"),
               ("source", "Source")]
STATE_WORDS = {"ok": "yes", "failed": "no - failed", "not_scanned": "not scanned yet"}


def _cell(v: Any) -> str:
    s = "" if v is None else str(v)
    # A cell starting with = + - @ would run as a formula in Excel.
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s


@router.get("/companies/export.csv")
def export_csv(request: Request) -> Response:
    """The same list as a CSV file to download: scannable companies first,
    then the ones whose fetch failed, with the reason (M17)."""
    with open_store(request) as store:
        view = company_health_view(store)
    buf = io.StringIO()
    out = csv.writer(buf)
    out.writerow([title for _, title in CSV_COLUMNS])
    for row in view["items"]:
        row = dict(row, state=STATE_WORDS[row["state"]])
        out.writerow([_cell(row[key]) for key, _ in CSV_COLUMNS])
    name = f"jobscraper-companies-{_dt.date.today().isoformat()}.csv"
    return Response(content="\ufeff" + buf.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.delete("/companies/list")
def clear_list(request: Request, cfg: Config = Depends(get_config)) -> dict[str, Any]:
    """Forget the list. The companies it added stop being watched at once."""
    with open_store(request) as store:
        before = _found_keys(store)
        store.replace_company_list([])
        _unwatch_dropped(cfg, store, before)
        return companies_view(cfg, store)


def _found_keys(store) -> set:
    return {r["company_key"] for r in store.company_list() if r["status"] == "found"}


def _unwatch_dropped(cfg: Config, store, before: set) -> None:
    """A company found from the list and now dropped from it stops being
    watched now, not at the next run - the Settings tab's company count and the
    next run's due list both see it at once. An unreadable watchlist.yaml
    leaves it to the next run's sync, which reports the error."""
    if before - _found_keys(store):
        try:
            store.sync_watchlist(watchlist.all_entries(cfg.watchlist_path, store))
        except watchlist.WatchlistError:
            pass


def _text(media: str, body: bytes) -> str:
    if media == DOCX_TYPE:
        if not body.startswith(b"PK\x03\x04"):
            raise HTTPException(status_code=415, detail="that is not a DOCX file")
        try:
            import docx                         # python-docx, also used for resumes
            doc = docx.Document(io.BytesIO(body))
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
            raise HTTPException(status_code=415, detail=f"that DOCX could not be read ({exc})")
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:                 # a list pasted as a table
            for row in table.rows:
                lines.extend(cell.text for cell in row.cells[:1])
        return "\n".join(lines)
    if body.startswith(b"PK\x03\x04") or b"\x00" in body[:1024]:
        raise HTTPException(status_code=415, detail="that is not a text file")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return body.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise HTTPException(status_code=415, detail="the text file could not be decoded")
