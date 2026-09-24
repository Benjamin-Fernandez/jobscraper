"""Synthetic resumes for the profile tests - never the real one.

The real resume lives at `data/resume.pdf`, is git-ignored, and the repo is
public. Tests therefore build their own: a one-page PDF and a DOCX for an
invented candidate, written into a temp directory at test time so no binary
fixture is ever committed.

The PDF is written by hand rather than with reportlab. A single page of
Helvetica text is about forty lines of PDF syntax, and hand-writing it keeps the
test suite free of a dependency the product itself does not need. pypdf - which
the product does need - reads it back, which is exactly the round trip the
ingest tests exercise.

    python tests/fixtures/make_resume_fixture.py out/dir    writes both files
"""
from __future__ import annotations

import sys
from pathlib import Path

# An invented candidate. Every line is plain ASCII so the PDF needs no font
# encoding beyond Helvetica's standard one.
LINES = [
    "Alex Tan",
    "Software Engineer - Singapore",
    "Education: BSc Computer Science, graduating 2026-08",
    "Skills: Python, Java, Spring Boot, FastAPI, Kubernetes, PostgreSQL, Kafka",
    "Experience: Backend intern, built REST APIs in Python and deployed on Docker",
    "Projects: CI/CD pipeline with GitHub Actions and ArgoCD on AWS",
]


def _pdf_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(path: str | Path, lines: list[str] | None = None) -> Path:
    """Write a valid one-page PDF whose text layer is `lines`, one per line."""
    path = Path(path)
    lines = LINES if lines is None else lines
    ops = ["BT", "/F1 11 Tf", "14 TL", "72 720 Td"]
    for line in lines:
        ops.append(f"({_pdf_escape(line)}) Tj T*")
    ops.append("ET")
    content = "\n".join(ops).encode("ascii")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
        + content + b"\nendstream",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n").encode()

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bytes(out))
    return path


def make_docx(path: str | Path, lines: list[str] | None = None) -> Path:
    """Write a DOCX with one paragraph per line, plus one table cell.

    The table is there because real resumes lay out skills in tables, and a
    paragraph-only extractor silently loses them.
    """
    import docx

    path = Path(path)
    lines = LINES if lines is None else lines
    doc = docx.Document()
    for line in lines:
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Tools: Terraform, Ansible"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    print(make_pdf(target / "resume.pdf"))
    print(make_docx(target / "resume.docx"))
