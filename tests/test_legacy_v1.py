"""v1 behaviour still under test, with a scheduled end date.

Nothing here is v2 design. Each group guards behaviour that still runs today and
has a named task that retires or rewrites it. The file shrinks to nothing and is
deleted; it must never grow.

  cursor tests (3)   -> M3-T3 ports them to tests/test_scheduler.py as staleness
                        equivalents, then deletes them here. cursor.py is RETIRE
                        (D-9) but still live until M9-T1.
  applied tests (6)  -> M6-T2 rewrites them against the new /api/applications
                        endpoint in tests/test_web.py. The BEHAVIOURS they assert
                        are v2 requirements - a date set once and kept, a status
                        never walked backwards, the user's own columns left alone
                        - only the transport changes.

Deleted in M0-T3, recorded here so the loss is deliberate and auditable:
  test_html_view_renders_the_live_run_path_and_the_tracker_path
  test_cards_from_tracker_survives_a_missing_file
Both tested output.py, which M9-T1 retires. v2 generates no HTML files at all
(PRD section 5), so there is nothing for them to assert.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import tempfile
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook, load_workbook

from jobscraper import cursor as C
from jobscraper.models import RawJob
from jobscraper.output import sync_applied
from jobscraper.store import Store

T0 = datetime(2026, 9, 14, 9, 0, tzinfo=timezone.utc)
PROG = 14          # program columns; Status is 15 (O), Date Applied is 16 (P)

def _store(n=25):
    db = Path(tempfile.mkdtemp()) / "t.db"
    st = Store(db)
    st.sync_companies([{"ordinal": i, "name": f"C{i}", "tier": "T3",
                        "category": "Fintech", "careers_url": "https://x/careers",
                        "role_type_hint": "SWE"} for i in range(1, n + 1)])
    return st

def _advance(st, comps, now, force=False):
    p = C.plan_batch(st, comps, 10, 14, force=force, now=now)
    if p.blocked:
        return "BLOCKED"
    run = st.start_run(p.cycle.cycle_id, p.start_ordinal, p.end_ordinal)
    C.advance(st, p, run)
    return f"{p.start_ordinal}-{p.end_ordinal}@c{p.cycle.cycle_id}"

def _tracker(tmp, rows):
    """A minimal stand-in tracker: job_id in A, Status in O, Date in P."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append(["job_id"] + [""] * (PROG - 1) + ["Status", "Date Applied",
                                                "Resume Version", "Referral",
                                                "Notes", "Outcome"])
    for job_id, status, date, notes in rows:
        ws.append([job_id] + [""] * (PROG - 1) + [status, date, "", "", notes, ""])
    path = Path(tmp) / "tracker.xlsx"
    wb.save(path)
    return path

def test_cursor_advances_ten_at_a_time():
    st = _store(25)
    comps = st.all_companies()
    assert _advance(st, comps, T0) == "1-10@c1"
    assert _advance(st, comps, T0 + timedelta(days=1)) == "11-20@c1"
    assert _advance(st, comps, T0 + timedelta(days=1, hours=2)) == "21-25@c1"

def test_cycle_blocked_until_fourteen_days():
    st = _store(25)
    comps = st.all_companies()
    for d in (0, 1, 1):
        _advance(st, comps, T0 + timedelta(days=d))
    assert _advance(st, comps, T0 + timedelta(days=2)) == "BLOCKED"
    assert _advance(st, comps, T0 + timedelta(days=13, hours=23)) == "BLOCKED"
    assert _advance(st, comps, T0 + timedelta(days=14, minutes=1)) == "1-10@c2"

def test_force_overrides_the_wait():
    st = _store(25)
    comps = st.all_companies()
    for d in (0, 1, 1):
        _advance(st, comps, T0 + timedelta(days=d))
    assert _advance(st, comps, T0 + timedelta(days=2), force=True) == "1-10@c2"

def test_applied_tick_writes_status_and_date():
    tmp = tempfile.mkdtemp()
    st = _store(1)
    path = _tracker(tmp, [("j1", "To Apply", None, "")])
    st.set_applied("j1", True, "SWE", "Acme", "https://e/1")
    sync_applied(path, st)
    ws = load_workbook(path)["Applications"]
    assert ws.cell(row=2, column=15).value == "Applied"
    assert ws.cell(row=2, column=16).value == st.get_application("j1")["applied_at"]

def test_applied_never_walks_back_a_status_you_advanced():
    tmp = tempfile.mkdtemp()
    st = _store(1)
    path = _tracker(tmp, [("j1", "Interview", "2026-09-01", "")])
    st.set_applied("j1", True, "SWE", "Acme", "https://e/1")
    assert sync_applied(path, st) == 0          # nothing touched at all
    ws = load_workbook(path)["Applications"]
    assert ws.cell(row=2, column=15).value == "Interview"
    assert ws.cell(row=2, column=16).value == "2026-09-01"

def test_applied_sync_leaves_your_own_columns_alone():
    tmp = tempfile.mkdtemp()
    st = _store(1)
    path = _tracker(tmp, [("j1", "To Apply", None, "referred by a friend")])
    st.set_applied("j1", True, "SWE", "Acme", "https://e/1")
    sync_applied(path, st)
    ws = load_workbook(path)["Applications"]
    assert ws.cell(row=2, column=19).value == "referred by a friend"

def test_untick_reverts_only_the_applied_status():
    tmp = tempfile.mkdtemp()
    st = _store(1)
    path = _tracker(tmp, [("j1", "To Apply", None, "")])
    st.set_applied("j1", True, "SWE", "Acme", "https://e/1")
    sync_applied(path, st)
    st.set_applied("j1", False)
    sync_applied(path, st)
    ws = load_workbook(path)["Applications"]
    assert ws.cell(row=2, column=15).value == "To Apply"
    assert ws.cell(row=2, column=16).value is None

def test_applied_date_is_set_once_and_kept():
    st = _store(1)
    first = st.set_applied("j1", True, "SWE", "Acme", "https://e/1")["applied_at"]
    again = st.set_applied("j1", True, "SWE", "Acme", "https://e/1")["applied_at"]
    assert first == again

def test_applied_sheet_logs_role_date_and_link():
    tmp = tempfile.mkdtemp()
    st = _store(1)
    path = _tracker(tmp, [("j1", "To Apply", None, "")])
    st.set_applied("j1", True, "Software Engineer", "Acme", "https://e/1")
    sync_applied(path, st)
    wb = load_workbook(path)
    assert "Applied" in wb.sheetnames
    row = list(wb["Applied"].iter_rows(min_row=2, values_only=True))[0]
    assert row[1] == "Acme" and row[3] == "Software Engineer"
    assert wb["Applied"].cell(row=2, column=6).hyperlink.target == "https://e/1"
