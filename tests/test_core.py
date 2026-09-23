"""Offline regression tests. No network.

    python -m pytest tests -q        (or)     python tests/test_core.py
"""
from __future__ import annotations

import json
import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from jobscraper import backends as B                    # noqa: E402
from jobscraper import cursor as C                      # noqa: E402
from jobscraper import discovery as D                   # noqa: E402
from jobscraper import review as R                      # noqa: E402
from jobscraper.config import load_config, load_profile  # noqa: E402
from jobscraper.llm import _key                         # noqa: E402
from jobscraper.matching import (location_verdict,      # noqa: E402
                                 min_years_required, stage_a)
from jobscraper.models import Candidate, RawJob         # noqa: E402
from jobscraper.models import Company, ScoreBreakdown   # noqa: E402
from jobscraper.output import (cards_from_tracker,      # noqa: E402
                               sync_applied, write_html_view)
from openpyxl import Workbook, load_workbook            # noqa: E402
from jobscraper.store import Store                      # noqa: E402

T0 = datetime(2026, 9, 14, 9, 0, tzinfo=timezone.utc)


def _store(n=25):
    db = Path(tempfile.mkdtemp()) / "t.db"
    st = Store(db)
    st.sync_companies([{"ordinal": i, "name": f"C{i}", "tier": "T3",
                        "category": "Fintech", "careers_url": "https://x/careers",
                        "role_type_hint": "SWE"} for i in range(1, n + 1)])
    return st


def _advance(st, comps, now, force=False):
    p = C.plan_batch(st, comps, 10, 14, force=force, now=now)
    if p.blocked:
        return "BLOCKED"
    run = st.start_run(p.cycle.cycle_id, p.start_ordinal, p.end_ordinal)
    C.advance(st, p, run)
    return f"{p.start_ordinal}-{p.end_ordinal}@c{p.cycle.cycle_id}"


def test_cursor_advances_ten_at_a_time():
    st = _store(25)
    comps = st.all_companies()
    assert _advance(st, comps, T0) == "1-10@c1"
    assert _advance(st, comps, T0 + timedelta(days=1)) == "11-20@c1"
    assert _advance(st, comps, T0 + timedelta(days=1, hours=2)) == "21-25@c1"


def test_cycle_blocked_until_fourteen_days():
    st = _store(25)
    comps = st.all_companies()
    for d in (0, 1, 1):
        _advance(st, comps, T0 + timedelta(days=d))
    assert _advance(st, comps, T0 + timedelta(days=2)) == "BLOCKED"
    assert _advance(st, comps, T0 + timedelta(days=13, hours=23)) == "BLOCKED"
    assert _advance(st, comps, T0 + timedelta(days=14, minutes=1)) == "1-10@c2"


def test_force_overrides_the_wait():
    st = _store(25)
    comps = st.all_companies()
    for d in (0, 1, 1):
        _advance(st, comps, T0 + timedelta(days=d))
    assert _advance(st, comps, T0 + timedelta(days=2), force=True) == "1-10@c2"


def test_failed_fetch_must_not_close_jobs():
    """The trap: a timeout must never mass-close a company's postings."""
    st = _store(3)
    cid = st.all_companies()[0].id
    raw = RawJob(external_id="1", title="Software Engineer", url="https://x/1",
                 location="Singapore")
    jid = raw.job_id(cid, "greenhouse")
    st.upsert_job(jid, cid, raw, 1, True)
    st.commit()

    # A SUCCESSFUL fetch that omits the job closes it...
    assert st.close_missing(cid, set()) == 1
    # ...and close_missing is the only path that closes anything; runner.py calls
    # it exclusively for outcomes where ok=True. Re-seeing the job reopens it.
    st.upsert_job(jid, cid, raw, 2, False)
    st.commit()
    row = st.conn.execute(
        "SELECT closed_at FROM jobs WHERE job_id=?", (jid,)).fetchone()
    assert row[0] is None


def test_location_is_default_deny():
    p = load_profile()
    assert location_verdict("Singapore", p) == "singapore"
    assert location_verdict("Singapore, Hong Kong", p) == "singapore"
    for elsewhere in ("Zug, Switzerland", "Aarhus, Central Denmark Region",
                      "Europe", "New York", "London, UK", "Bengaluru"):
        assert location_verdict(elsewhere, p) == "elsewhere", elsewhere
    for vague in ("", "APAC", "Remote", "All Offices", "Multiple Locations"):
        assert location_verdict(vague, p) == "ambiguous", vague


def test_remote_pinned_to_a_region_is_elsewhere():
    """Remote is only vague while it is unqualified.

    A desk we cannot legally occupy is out of reach however remote it is, so a
    region-locked remote posting is rejected like any other foreign location.
    The last case names a city on no deny list: the rule has to fail closed on
    residue, not on enumeration, or the next unlisted city leaks through again.
    """
    p = load_profile()
    for pinned in ("US Remote", "US-Remote", "US - Remote", "US Remote National",
                   "Remote (US)", "Remote-Canada", "EMEA Remote",
                   "SF, NY, Remote", "NYC, SF, Chi, Remote",
                   "Chicago, Atlanta, Remote", "Toronto, Remote-Canada",
                   "Remote - Munich"):
        assert location_verdict(pinned, p) == "elsewhere", pinned

    # Genuinely borderless remote still reaches the model.
    for open_ended in ("Remote (Worldwide)", "Work from anywhere",
                       "Fully Remote", "Remote - APAC"):
        assert location_verdict(open_ended, p) == "ambiguous", open_ended

    # A Singapore desk is not collateral damage when remote is also offered.
    assert location_verdict("Singapore / Remote", p) == "singapore"
    assert location_verdict("Remote - Singapore", p) == "singapore"


def test_stage_a_rejects_and_keeps_the_right_things():
    p = load_profile()
    st = _store(1)
    co = st.all_companies()[0]

    def judge(title, loc="Singapore", desc=""):
        c = Candidate(job_id="x", company=co,
                      raw=RawJob("1", title, "https://x", loc, description=desc))
        return stage_a(c, p), c.filter_reason

    assert judge("Software Engineer")[0]
    assert judge("Graduate Technology Analyst")[0]
    assert judge("Quantitative Developer")[0]
    assert not judge("Senior Software Engineer")[0]
    assert not judge("Software Engineer Intern")[0]
    assert not judge("HR Generalist")[0]
    assert not judge("Software Engineer", loc="New York")[0]
    assert not judge("Software Engineer", desc="Requires 7+ years experience")[0]
    # A vague location survives rather than being guessed away.
    assert judge("Software Engineer", loc="APAC")[0]


def test_min_years_takes_the_lowest_stated():
    assert min_years_required("2+ years preferred, 5 years ideal") == 2
    assert min_years_required("no numbers here") is None


# ---------------------------------------------------------------- applied ---

PROG = 14          # program columns; Status is 15 (O), Date Applied is 16 (P)


def _tracker(tmp, rows):
    """A minimal stand-in tracker: job_id in A, Status in O, Date in P."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append(["job_id"] + [""] * (PROG - 1) + ["Status", "Date Applied",
                                                "Resume Version", "Referral",
                                                "Notes", "Outcome"])
    for job_id, status, date, notes in rows:
        ws.append([job_id] + [""] * (PROG - 1) + [status, date, "", "", notes, ""])
    path = Path(tmp) / "tracker.xlsx"
    wb.save(path)
    return path


def test_applied_tick_writes_status_and_date():
    tmp = tempfile.mkdtemp()
    st = _store(1)
    path = _tracker(tmp, [("j1", "To Apply", None, "")])
    st.set_applied("j1", True, "SWE", "Acme", "https://e/1")
    sync_applied(path, st)
    ws = load_workbook(path)["Applications"]
    assert ws.cell(row=2, column=15).value == "Applied"
    assert ws.cell(row=2, column=16).value == st.get_application("j1")["applied_at"]


def test_applied_never_walks_back_a_status_you_advanced():
    tmp = tempfile.mkdtemp()
    st = _store(1)
    path = _tracker(tmp, [("j1", "Interview", "2026-09-01", "")])
    st.set_applied("j1", True, "SWE", "Acme", "https://e/1")
    assert sync_applied(path, st) == 0          # nothing touched at all
    ws = load_workbook(path)["Applications"]
    assert ws.cell(row=2, column=15).value == "Interview"
    assert ws.cell(row=2, column=16).value == "2026-09-01"


def test_applied_sync_leaves_your_own_columns_alone():
    tmp = tempfile.mkdtemp()
    st = _store(1)
    path = _tracker(tmp, [("j1", "To Apply", None, "referred by a friend")])
    st.set_applied("j1", True, "SWE", "Acme", "https://e/1")
    sync_applied(path, st)
    ws = load_workbook(path)["Applications"]
    assert ws.cell(row=2, column=19).value == "referred by a friend"


def test_untick_reverts_only_the_applied_status():
    tmp = tempfile.mkdtemp()
    st = _store(1)
    path = _tracker(tmp, [("j1", "To Apply", None, "")])
    st.set_applied("j1", True, "SWE", "Acme", "https://e/1")
    sync_applied(path, st)
    st.set_applied("j1", False)
    sync_applied(path, st)
    ws = load_workbook(path)["Applications"]
    assert ws.cell(row=2, column=15).value == "To Apply"
    assert ws.cell(row=2, column=16).value is None


def test_applied_date_is_set_once_and_kept():
    st = _store(1)
    first = st.set_applied("j1", True, "SWE", "Acme", "https://e/1")["applied_at"]
    again = st.set_applied("j1", True, "SWE", "Acme", "https://e/1")["applied_at"]
    assert first == again


def test_applied_sheet_logs_role_date_and_link():
    tmp = tempfile.mkdtemp()
    st = _store(1)
    path = _tracker(tmp, [("j1", "To Apply", None, "")])
    st.set_applied("j1", True, "Software Engineer", "Acme", "https://e/1")
    sync_applied(path, st)
    wb = load_workbook(path)
    assert "Applied" in wb.sheetnames
    row = list(wb["Applied"].iter_rows(min_row=2, values_only=True))[0]
    assert row[1] == "Acme" and row[3] == "Software Engineer"
    assert wb["Applied"].cell(row=2, column=6).hyperlink.target == "https://e/1"


def test_html_view_renders_the_live_run_path_and_the_tracker_path():
    """`run` passes Candidates; `view` passes tracker dicts. Both must render."""
    tmp = Path(tempfile.mkdtemp())
    co = Company(id=1, ordinal=1, name="Acme", tier="T1", category="Fintech",
                 careers_url="https://acme/careers")
    raw = RawJob(external_id="1", title="Software Engineer", location="Singapore",
                 url="https://acme/jobs/1")
    cand = Candidate(job_id="j1", company=co, raw=raw,
                     score=ScoreBreakdown(total=72.0, skill_overlap=40,
                                          title_affinity=60))
    live = tmp / "live.html"
    write_html_view(live, [cand], 1, "companies 1-10")
    doc = live.read_text(encoding="utf-8")
    assert 'data-job="j1"' in doc
    assert 'type="checkbox"' in doc
    assert "Software Engineer" in doc

    empty = tmp / "empty.html"
    write_html_view(empty, [], 1, "companies 1-10")
    assert "No new matches" in empty.read_text(encoding="utf-8")

    from_dicts = tmp / "dicts.html"
    write_html_view(from_dicts, [{"job_id": "j2", "title": "Backend Engineer",
                                  "company": "Acme", "tier": "T2",
                                  "location": "Singapore", "url": "https://a/2",
                                  "score": 61.0, "verdict": "strong",
                                  "why": "because"}], 0, "all", heading="All matches")
    doc2 = from_dicts.read_text(encoding="utf-8")
    assert "Backend Engineer" in doc2 and "All matches" in doc2


def test_cards_from_tracker_survives_a_missing_file():
    assert cards_from_tracker(Path(tempfile.mkdtemp()) / "nope.xlsx") == []


# -------------------------------------------------------------- discovery ---

def test_aggregator_urls_never_yield_a_slug():
    """Mathrix once became LinkedIn because its careers URL was a LinkedIn page."""
    agg = "https://www.linkedin.com/company/mathrix/jobs/"
    assert D.is_aggregator(agg)
    assert D.domain_slug(agg) is None
    assert D.domain_slug("https://www.hudsonrivertrading.com/careers/") ==         "hudsonrivertrading"


def test_board_must_corroborate_the_company_name():
    # anyone can register google.recruitee.com
    assert not D.corroborates("Meta", "Facebook Data", "meta")
    assert not D.corroborates("Mathrix", "LinkedIn", "linkedin")
    assert D.corroborates("Jane Street", "Jane Street", "janestreet")
    assert D.corroborates("IMC Trading", "IMC", "imc")


def test_slug_alone_corroborates_when_the_ats_reports_no_name():
    assert D.corroborates("Virtu Financial", None, "virtu")
    assert not D.corroborates("Mathrix", None, "linkedin")


def test_demo_boards_are_rejected():
    assert D._all_samples(["Senior Marketer (Sample)"])
    assert D._all_samples(["Test Job", "Your first job"])
    assert not D._all_samples(["Software Engineer", "Sample Analyst"])
    assert not D._all_samples([])


def test_model_ids_map_onto_cli_aliases():
    assert B.cli_model_alias("claude-sonnet-5") == "sonnet"
    assert B.cli_model_alias("claude-haiku-4-5-20251001") == "haiku"
    assert B.cli_model_alias("claude-opus-5") == "opus"
    # Anything unrecognised is handed to the CLI untouched.
    assert B.cli_model_alias("some-future-model") == "some-future-model"


def test_backend_choice_follows_config():
    assert B.build({"enable_llm": False, "backend": "cli"}).name == "off"
    assert B.build({"enable_llm": True, "backend": "off"}).name == "off"
    assert B.build({"enable_llm": True, "backend": "cli"}).name == "cli"
    bad = B.build({"enable_llm": True, "backend": "nonsense"})
    assert bad.name == "off" and "nonsense" in bad.unavailable_reason


def test_backend_says_so_when_the_cli_is_not_installed():
    # An empty `binary` means auto-detect, so hide `claude` from PATH lookup to
    # reach the not-installed path on a machine that does have it.
    real_which = B.shutil.which
    B.shutil.which = lambda _name: None
    try:
        b = B.ClaudeCliBackend(binary="")
        assert not b.available and "not on PATH" in b.unavailable_reason
        try:
            b.complete("claude-haiku-4-5", "sys", "user")
            raise AssertionError("a missing CLI must refuse, not pretend")
        except B.BackendError:
            pass
    finally:
        B.shutil.which = real_which


def test_a_broken_cli_path_fails_loudly_rather_than_hanging():
    b = B.ClaudeCliBackend(binary=str(Path(tempfile.mkdtemp()) / "no-such-claude"))
    assert b.available, "a configured path is taken at face value until used"
    try:
        b.complete("claude-haiku-4-5", "sys", "user")
        raise AssertionError("a bad binary path must raise")
    except B.BackendError as exc:
        assert "could not start" in str(exc)


def test_cli_envelope_survives_a_banner_line():
    assert B._envelope('{"result": "hi"}')["result"] == "hi"
    noisy = 'Some notice line\n{"result": "hi", "is_error": false}'
    assert B._envelope(noisy)["result"] == "hi"
    for junk in ("", "not json at all"):
        try:
            B._envelope(junk)
            raise AssertionError(f"{junk!r} should not parse")
        except B.BackendError:
            pass


def _scored_job(st, cid, jid, title, score, desc="Backend role in Singapore."):
    raw = RawJob(external_id=jid, title=title, url=f"https://x/{jid}",
                 location="Singapore", description=desc)
    st.upsert_job(jid, cid, raw, 1, True)
    st.commit()
    cand = Candidate(job_id=jid, company=st.get_company(cid), raw=raw)
    cand.score = ScoreBreakdown(total=score)
    st.save_score(cand, 1)
    return raw


def _review_setup(tmp):
    """A store plus a Config/Profile pointing at a throwaway output dir."""
    st = _store(3)
    cfg = load_config()
    cfg.raw["paths"]["output_dir"] = str(tmp)
    return st, cfg, load_profile()


def test_review_queue_holds_only_unjudged_in_band_postings():
    tmp = Path(tempfile.mkdtemp())
    st, cfg, prof = _review_setup(tmp)
    lo, hi = prof.thresholds["llm_band"]
    _scored_job(st, 1, "inband", "Graduate Software Engineer", (lo + hi) / 2)
    _scored_job(st, 1, "toolow", "Cleaner", float(lo) - 5)
    _scored_job(st, 1, "toohigh", "Perfect Match", float(hi) + 5)

    path, n = R.export_queue(cfg, prof, st)
    assert n == 1, f"expected only the in-band posting, got {n}"
    queued = json.loads(path.read_text(encoding="utf-8"))
    assert [j["id"] for j in queued["jobs"]] == ["inband"]
    # The queue has to carry everything a reviewer needs to decide.
    job = queued["jobs"][0]
    for field in ("id", "company", "title", "location", "score", "description"):
        assert field in job, f"queue entry is missing {field}"
    assert queued["candidate_profile"], "reviewer needs the candidate profile"


def test_applied_verdicts_are_cached_and_never_re_queued():
    tmp = Path(tempfile.mkdtemp())
    st, cfg, prof = _review_setup(tmp)
    lo, hi = prof.thresholds["llm_band"]
    mid = (float(lo) + float(hi)) / 2
    _scored_job(st, 1, "keepme", "Graduate Backend Engineer", mid)
    _scored_job(st, 2, "dropme", "Senior Staff Engineer", mid)

    _path, n = R.export_queue(cfg, prof, st)
    assert n == 2

    (tmp / R.VERDICTS_NAME).write_text(json.dumps({"results": [
        {"id": "keepme", "verdict": "strong", "confidence": 0.9,
         "reason": "grad role", "concerns": []},
        {"id": "dropme", "verdict": "reject", "confidence": 0.9,
         "reason": "too senior", "concerns": []},
    ]}), encoding="utf-8")

    stats = R.apply_verdicts(cfg, prof, st)
    assert stats["applied"] == 2, stats
    assert stats["tally"] == {"strong": 1, "reject": 1}, stats["tally"]
    # Only the strong one is exportable; the rejected one must not reach the tracker.
    assert stats["exported"] == 1, stats

    _path, n = R.export_queue(cfg, prof, st)
    assert n == 0, "a judged posting must not come round again"


def test_apply_ignores_unknown_ids_and_junk_verdicts():
    tmp = Path(tempfile.mkdtemp())
    st, cfg, prof = _review_setup(tmp)
    lo, hi = prof.thresholds["llm_band"]
    _scored_job(st, 1, "real", "Graduate Backend Engineer",
                (float(lo) + float(hi)) / 2)

    (tmp / R.VERDICTS_NAME).write_text(json.dumps({"results": [
        {"id": "real", "verdict": "STRONG", "confidence": 5, "reason": "x"},
        {"id": "ghost", "verdict": "strong", "confidence": 0.5, "reason": "x"},
        {"id": "real", "verdict": "maybe-ish", "confidence": 0.5, "reason": "x"},
    ]}), encoding="utf-8")

    stats = R.apply_verdicts(cfg, prof, st)
    assert stats["applied"] == 1, stats
    assert stats["unknown_id"] == 1 and stats["bad_verdict"] == 1, stats
    # Upper-cased verdict normalised, out-of-range confidence clamped.
    jd_hash = RawJob(external_id="real", title="Graduate Backend Engineer",
                     url="https://x/real", location="Singapore",
                     description="Backend role in Singapore.").jd_hash()
    row = st.get_cached(_key("real", jd_hash, 1, "c2"))
    assert row and row["verdict"] == "strong" and row["confidence"] == 1.0, row


def test_missing_verdicts_file_is_a_clear_error():
    tmp = Path(tempfile.mkdtemp())
    st, cfg, prof = _review_setup(tmp)
    try:
        R.apply_verdicts(cfg, prof, st)
        raise AssertionError("a missing verdicts file must not pass silently")
    except SystemExit as exc:
        assert "review --export" in str(exc)


if __name__ == "__main__":
    fns = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for fn in fns:
        try:
            fn()
            print(f"  PASS  {fn.__name__}")
        except AssertionError as exc:
            failed += 1
            print(f"  FAIL  {fn.__name__}: {exc}")
    print(f"\n{len(fns) - failed}/{len(fns)} passed")
    raise SystemExit(1 if failed else 0)
